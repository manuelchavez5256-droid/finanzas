"""
Convierte la tabla "Registro" de Finanzas.xlsx en data/importado.json,
que la app ofrece cargar la primera vez que se abre.

Uso:
    python herramientas/importar_excel.py "C:/ruta/a/Finanzas.xlsx"
"""
import json
import sys
import uuid
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent


def limpiar(texto):
    return " ".join(str(texto or "").split())


def main(ruta_excel):
    # data_only=True: lee los valores ya calculados en lugar de las fórmulas
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    hoja = wb["Dashboard"]
    tabla = hoja.tables["Registro"]
    filas = hoja[tabla.ref]
    encabezados = [limpiar(c.value) for c in filas[0]]

    movimientos = []
    for fila in filas[1:]:
        datos = dict(zip(encabezados, (c.value for c in fila)))
        if not datos.get("Fecha"):
            continue
        movimientos.append({
            "id": str(uuid.uuid4()),
            "fecha": datos["Fecha"].strftime("%Y-%m-%d"),
            "tipo": limpiar(datos["Tipo"]),
            "categoria": limpiar(datos["Categoría"]),
            "concepto": limpiar(datos["Concepto"]),
            "presupuesto": float(datos["Presupuesto"] or 0),
            "monto": float(datos["Monto"] or 0),
        })

    # Categorías: las de la hoja "Lista Validadas" más cualquiera usada en el registro
    listas = wb["Lista Validadas"]
    categorias = {"Ingreso": [], "Egreso": []}
    for fila in listas.iter_rows(min_row=2, values_only=True):
        if fila[2] and limpiar(fila[2]) not in categorias["Ingreso"]:
            categorias["Ingreso"].append(limpiar(fila[2]))
        if fila[4] and limpiar(fila[4]) not in categorias["Egreso"]:
            categorias["Egreso"].append(limpiar(fila[4]))
    for m in movimientos:
        if m["categoria"] not in categorias[m["tipo"]]:
            categorias[m["tipo"]].append(m["categoria"])

    salida = RAIZ / "data" / "importado.json"
    salida.write_text(
        json.dumps({"movimientos": movimientos, "categorias": categorias}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    print(f"{len(movimientos)} movimientos exportados a {salida}")
    print("Categorías:", categorias)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
